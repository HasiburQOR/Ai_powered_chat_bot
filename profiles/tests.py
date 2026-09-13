"""Travel-profile capture: numbering, LLM extraction, engine intro bubble."""
import datetime as dt
import html
import re
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from bot.engine import handle_inbound_message
from conversations.models import Conversation, Customer, Message
from knowledge.models import BOT_SETTINGS_CACHE_KEY, BotSettings, Rule
from llm.adapters import LLMProviderError
from llm.models import LLMConfig
from platforms.models import Channel
from profiles.extraction import (
    EXTRACTION_TEMPERATURE,
    apply_fields,
    extract_fields,
    parse_extraction_json,
)
from profiles.models import ProfileNumberCounter, TravelProfile
from profiles.tasks import extract_profile_task


class ProfileNumberTests(TestCase):
    def test_numbers_increment_and_format(self):
        self.assertEqual(ProfileNumberCounter.next_profile_number(), "BP-000001")
        self.assertEqual(ProfileNumberCounter.next_profile_number(), "BP-000002")
        self.assertEqual(TravelProfile.objects.count(), 0)  # numbers are minted independently


class ExtractionTests(TestCase):
    def test_parse_tolerates_code_fences_and_prose(self):
        raw = 'Sure! ```json\n{"full_name": "Ravi"}\n``` hope that helps'
        self.assertEqual(parse_extraction_json(raw), {"full_name": "Ravi"})

    def test_parse_garbage_returns_empty(self):
        self.assertEqual(parse_extraction_json("no json here"), {})
        self.assertEqual(parse_extraction_json(""), {})

    def test_extract_fields_coerces_and_drops_invalid(self):
        def fake_llm(messages, config, **kwargs):
            return ('```json\n{"full_name": "Ravi", "trip_days": "12", "adults": 2, '
                    '"travel_date": "2026-10-12", "junk": "x"}\n```')

        fields = extract_fields("my name is Ravi, 12 days, 2 adults", send_fn=fake_llm)
        self.assertEqual(fields, {
            "full_name": "Ravi", "trip_days": 12, "adults": 2,
            "travel_date": dt.date(2026, 10, 12),
        })

    def test_extract_fields_bad_types_dropped(self):
        fields = extract_fields("x", send_fn=lambda m, c, **kw: '{"trip_days": "soon", "adults": null}')
        self.assertEqual(fields, {})

    def test_extraction_prompt_guards_against_live_llm_sloppiness(self):
        """Live regressions shipped by the flaky model: mangled phone digits
        (+5509324243 → 9242834034), "12 people, 2 children" → adults 7,
        "december 16" → 14 Dec, "expiry 2030" → hallucinated 15 Dec 2030.
        The system prompt must carry explicit counter-rules and actually be
        sent as the first message."""
        captured = {}

        def fake_llm(messages, config, **kwargs):
            captured["messages"] = messages
            return "{}"

        fields = extract_fields("anything", send_fn=fake_llm)
        self.assertEqual(fields, {})
        system_text = captured["messages"][0]["content"]
        self.assertIn("EXACTLY as the customer typed", system_text)
        self.assertIn("adults = total minus children", system_text)
        self.assertIn("never shift it", system_text)
        self.assertIn("year alone", system_text)

    def test_extract_fields_detects_travel_intent(self):
        fields = extract_fields(
            "we want to visit Dubai in December",
            send_fn=lambda m, c, **kw: '{"travel_intent": true}')
        self.assertEqual(fields, {"travel_intent": True})

    def test_blank_text_never_calls_llm(self):
        called = []

        def fake_llm(messages, config, **kwargs):
            called.append(1)
            return "{}"

        self.assertEqual(extract_fields("   ", send_fn=fake_llm), {})
        self.assertEqual(called, [])

    def test_no_active_config_returns_empty(self):
        self.assertEqual(extract_fields("hello"), {})


class ApplyFieldsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.channel = Channel.objects.create(name="WP", channel_type="wordpress")
        cls.customer = Customer.objects.create(channel=cls.channel, external_id="lead-1")

    def test_first_fill_creates_numbered_profile(self):
        profile = apply_fields(self.customer, {"full_name": "Ravi Kumar", "nationality": "Indian"})
        self.assertEqual(profile.profile_number, "BP-000001")
        self.assertFalse(profile.is_complete)
        self.assertEqual(profile.full_name, "Ravi Kumar")
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.display_name, "Ravi Kumar")  # backfilled

    def test_travel_intent_flag_is_latched(self):
        apply_fields(self.customer, {"travel_intent": True})
        profile = self.customer.travel_profile
        self.assertTrue(profile.travel_intent_detected)
        # A later message without intent never un-detects earlier interest.
        apply_fields(self.customer, {"travel_intent": False, "full_name": "Ravi"})
        profile.refresh_from_db()
        self.assertTrue(profile.travel_intent_detected)

    def test_second_apply_updates_same_profile_and_completes(self):
        apply_fields(self.customer, {"full_name": "Ravi"})
        profile = apply_fields(self.customer, {
            "whatsapp_number": "+911234567890", "nationality": "Indian",
            "residence_country": "India", "travel_date": dt.date(2026, 12, 1),
            "trip_days": 7, "adults": 3,
        })
        self.assertEqual(profile.profile_number, "BP-000001")
        self.assertTrue(profile.is_complete)
        self.assertIsNotNone(profile.completed_at)
        self.assertEqual(TravelProfile.objects.count(), 1)
        self.assertEqual(profile.total_travellers, 3)

    def test_children_count_toward_travellers(self):
        profile = apply_fields(self.customer, {"adults": 2, "children_ages": "5, 8"})
        self.assertEqual(profile.total_travellers, 4)

    def test_missing_fields_listing(self):
        profile = apply_fields(self.customer, {"full_name": "Ravi"})
        self.assertIn("WhatsApp number", profile.missing_fields())
        self.assertIn("Travel date", profile.missing_fields())


    def test_llm_failure_returns_none_for_retry(self):
        """An LLM transport failure is None (retry me), not {} (nothing
        found) — the live glm-5.3-flash empty-content bug hid behind {}."""
        def failing_llm(messages, config, **kwargs):
            raise RuntimeError("provider down")

        self.assertIsNone(extract_fields("hello", send_fn=failing_llm))

    def test_json_mode_400_falls_back_to_a_plain_retry(self):
        """Not every provider implements response_format; one rejects it with
        HTTP 400. The extraction must retry the identical call without the
        option instead of being lost."""
        calls = []

        def fake_llm(messages, config, **kwargs):
            calls.append(kwargs)
            if kwargs.get("response_format") is not None:
                raise LLMProviderError(
                    "HTTP 400 from stub-model: response_format not supported")
            return '{"full_name": "Ravi"}'

        fields = extract_fields("my name is Ravi", send_fn=fake_llm)
        self.assertEqual(fields, {"full_name": "Ravi"})
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[1].get("response_format"), None)
        self.assertEqual(calls[1].get("temperature"), EXTRACTION_TEMPERATURE)


class ExtractionTaskTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.channel = Channel.objects.create(name="WP", channel_type="wordpress")
        cls.customer = Customer.objects.create(channel=cls.channel, external_id="lead-9")

    def test_task_merges_extracted_fields_into_profile(self):
        with mock.patch("profiles.extraction.LLMConfig") as fake_config_cls, \
                mock.patch("profiles.extraction.get_adapter") as fake_get_adapter:
            fake_config_cls.objects.filter.return_value.first.return_value = object()
            fake_get_adapter.return_value.send = lambda messages, config, **kwargs: \
                '{"full_name": "Ravi", "travel_date": "2026-10-12", "travel_intent": true}'
            extract_profile_task(str(self.customer.pk), "I am Ravi, travelling in October")

        profile = self.customer.travel_profile
        self.assertEqual(profile.full_name, "Ravi")
        self.assertEqual(profile.travel_date, dt.date(2026, 10, 12))
        self.assertTrue(profile.travel_intent_detected)
        self.assertEqual(profile.profile_number, "BP-000001")

    def test_task_ignores_unknown_customer_and_blank_text(self):
        extract_profile_task("00000000-0000-0000-0000-000000000000", "hi")
        extract_profile_task(str(self.customer.pk), "   ")
        self.assertFalse(TravelProfile.objects.exists())

    def test_extraction_runs_over_the_conversation_window(self):
        """Details spread across turns must all reach the extractor: the live
        glm failure captured 2 of ~7 stated fields because each message was
        extracted in isolation."""
        conversation = Conversation.objects.create(
            customer=self.customer, last_message_at=timezone.now())
        for content in ("my name is Ravi Kumar", "whatsapp +911234567890",
                        "we are 2 adults"):
            Message.objects.create(
                conversation=conversation,
                sender_type=Message.SenderType.CUSTOMER, content=content)
        captured = {}

        def fake_llm(messages, config, **kwargs):
            captured["text"] = messages[1]["content"]
            return ('{"full_name": "Ravi Kumar", "whatsapp_number": '
                    '"+911234567890", "adults": 2}')

        with mock.patch("profiles.extraction.LLMConfig") as fake_config_cls, \
                mock.patch("profiles.extraction.get_adapter") as fake_get_adapter:
            fake_config_cls.objects.filter.return_value.first.return_value = object()
            fake_get_adapter.return_value.send = fake_llm
            extract_profile_task(str(self.customer.pk), "we are 2 adults")

        self.assertIn("my name is Ravi Kumar", captured["text"])
        self.assertIn("whatsapp +911234567890", captured["text"])
        profile = self.customer.travel_profile
        self.assertEqual(profile.full_name, "Ravi Kumar")
        self.assertEqual(profile.whatsapp_number, "+911234567890")
        self.assertEqual(profile.adults, 2)

    def test_extraction_window_is_a_labelled_transcript_of_both_sides(self):
        """Bot lines must reach the extractor too, labelled: a bare
        '01712345678' is uninterpretable alone, obvious right after the bot
        asks for the WhatsApp number. The call must also be cold + JSON
        mode."""
        conversation = Conversation.objects.create(
            customer=self.customer, last_message_at=timezone.now())
        pairs = [
            (Message.SenderType.BOT, "Sure! Which WhatsApp number reaches you best?"),
            (Message.SenderType.CUSTOMER, "01712345678"),
            (Message.SenderType.BOT, "Got it. How many people are travelling?"),
            (Message.SenderType.CUSTOMER, "4 adults"),
        ]
        for sender, content in pairs:
            Message.objects.create(
                conversation=conversation, sender_type=sender, content=content)
        captured = {}

        def fake_llm(messages, config, **kwargs):
            captured["text"] = messages[1]["content"]
            captured["kwargs"] = kwargs
            return '{"whatsapp_number": "01712345678", "adults": 4}'

        with mock.patch("profiles.extraction.LLMConfig") as fake_config_cls, \
                mock.patch("profiles.extraction.get_adapter") as fake_get_adapter:
            fake_config_cls.objects.filter.return_value.first.return_value = object()
            fake_get_adapter.return_value.send = fake_llm
            extract_profile_task(str(self.customer.pk), "4 adults")

        text = captured["text"]
        self.assertIn("[Bot] Sure! Which WhatsApp number", text)
        self.assertIn("[Customer] 01712345678", text)
        self.assertEqual(captured["kwargs"].get("temperature"),
                         EXTRACTION_TEMPERATURE)
        self.assertEqual(captured["kwargs"].get("response_format"),
                         {"type": "json_object"})

    def test_transient_llm_failure_is_retried_not_dropped(self):
        """The live glm-5.3-flash empty-content failure: extract_fields
        returned None and the task silently "succeeded", losing the visitor's
        details forever. It must raise Retry instead."""
        from celery.exceptions import Retry

        with mock.patch("profiles.extraction.extract_fields",
                        return_value=None), \
                mock.patch.object(extract_profile_task, "retry",
                                  side_effect=Retry()) as fake_retry:
            with self.assertRaises(Retry):
                extract_profile_task(str(self.customer.pk), "hello")
            fake_retry.assert_called_once()

    def test_extraction_gives_up_loudly_after_max_retries(self):
        """Retries exhausted → log a warning and return: never crash the
        caller, never pretend success."""
        from celery.exceptions import MaxRetriesExceededError

        with mock.patch("profiles.extraction.extract_fields",
                        return_value=None), \
                mock.patch.object(extract_profile_task, "retry",
                                  side_effect=MaxRetriesExceededError()):
            extract_profile_task(str(self.customer.pk), "hello")  # must not raise


class StubbedLLMMixin:
    """A working (stubbed) LLM so engine/widget plumbing tests exercise the
    success path: the engine no longer sends the fallback text on intent turns
    (the scripted questions carry the turn alone when the LLM fails), so
    "reply + questions = 2 bubbles" needs a reply that succeeds."""

    REPLY = "Sure — we have several Dubai packages available."

    @classmethod
    def setUpTestData(cls):
        LLMConfig.objects.create(
            name="Primary", provider="openai_compatible",
            model_name="stub-model", api_base_url="http://127.0.0.1:9/v1",
            is_active=True)

    def setUp(self):
        super().setUp()
        # BotSettings is cached for 30s — and the locmem cache, unlike the
        # DB, is NOT rolled back between tests. Without this, the class's own
        # "disable profile collection" test poisons the cache and every later
        # intent turn sees its scripted questions silently switched off.
        cache.delete(BOT_SETTINGS_CACHE_KEY)
        patcher = mock.patch("bot.engine.get_adapter")
        patcher.start().return_value.send.return_value = self.REPLY
        self.addCleanup(patcher.stop)


class EngineProfileIntroTests(StubbedLLMMixin, TestCase):
    """The scripted profile questions ride along as a second bot bubble — but
    only once, and only after the visitor shows travel intent. (The stubbed
    LLM keeps the "reply + questions" flow healthy; the LLM-failure path —
    questions only, no apology bubble — is covered by bot's
    EngineFallbackIntroTests and by WidgetSendTests below.)"""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()  # the mixin's active LLMConfig — without it
        # the engine sees config=None, skips the LLM loop, and the "reply +
        # scripted questions" flow degrades to questions only.
        cls.channel = Channel.objects.create(name="WP", channel_type="wordpress")
        cls.customer = Customer.objects.create(channel=cls.channel, external_id="lead-2")
        cls.conversation = Conversation.objects.create(
            customer=cls.customer, last_message_at=timezone.now())
        Rule.objects.create(name="greet", trigger_keywords=["hello"],
                            response_text="Hello! How can I help you?",
                            short_circuits_llm=True, priority=10)

    def test_plain_greeting_gets_reply_only(self):
        """No travel intent → no interrogation, and no lead row either."""
        outbounds = handle_inbound_message(self.conversation, "hello")
        self.assertEqual(len(outbounds), 1)
        self.assertEqual(outbounds[0].content, "Hello! How can I help you?")
        self.assertFalse(TravelProfile.objects.exists())

    def test_travel_intent_gets_reply_plus_questions_once(self):
        outbounds = handle_inbound_message(
            self.conversation, "hello, I want a 7 days Dubai package")
        self.assertEqual(len(outbounds), 2)
        self.assertIn("travelling", outbounds[1].content)
        self.assertTrue(all(m.sender_type == Message.SenderType.BOT for m in outbounds))
        profile = self.customer.travel_profile
        self.assertTrue(profile.travel_intent_detected)
        self.assertIsNotNone(profile.questions_sent_at)

        # Intent again, but the questions were already sent → reply only.
        again = handle_inbound_message(self.conversation, "maybe baku in spring too")
        self.assertEqual(len(again), 1)

    def test_llm_detected_intent_triggers_questions_without_keywords(self):
        """The extractor's travel_intent flag backstops non-English intent the
        keyword list can't see."""
        apply_fields(self.customer, {"travel_intent": True})
        outbounds = handle_inbound_message(self.conversation, "ok")
        self.assertEqual(len(outbounds), 2)

    def test_incomplete_profile_stays_silent_without_intent(self):
        apply_fields(self.customer, {"full_name": "Ravi"})
        outbounds = handle_inbound_message(self.conversation, "hello again")
        self.assertEqual(len(outbounds), 1)
        profile = TravelProfile.objects.get(customer=self.customer)
        self.assertIsNone(profile.questions_sent_at)

    def test_disabled_settings_skip_intro(self):
        settings = BotSettings.load()
        settings.profile_collection_enabled = False
        settings.save()
        outbounds = handle_inbound_message(
            self.conversation, "hello, I'd like a Dubai package")
        self.assertEqual(len(outbounds), 1)

    def test_complete_profile_skips_intro(self):
        apply_fields(self.customer, {
            "full_name": "Ravi", "whatsapp_number": "+911234567890",
            "nationality": "Indian", "residence_country": "India",
            "travel_date": dt.date(2026, 12, 1), "trip_days": 7, "adults": 2,
        })
        outbounds = handle_inbound_message(
            self.conversation, "dubai package for 2 please")
        self.assertEqual(len(outbounds), 1)

    def test_language_and_format_instructions_sent_to_llm(self):
        captured = {}

        class FakeAdapter:
            def send(self, messages, config):
                captured["messages"] = messages
                return "नमस्ते! मैं आपकी कैसे मदद कर सकता हूँ?"

        from llm.models import LLMConfig
        LLMConfig.objects.create(
            name="Test", provider="anthropic", api_key="k",
            model_name="claude-sonnet-4-6", system_prompt="Be helpful.", is_active=True)
        with mock.patch("bot.engine.get_adapter", return_value=FakeAdapter()):
            outbounds = handle_inbound_message(self.conversation, "namaste, kya haal?")
        system_text = " ".join(
            m["content"] for m in captured["messages"] if m["role"] == "system")
        self.assertIn("same language", system_text.lower())
        self.assertIn("**bold**", system_text)
        self.assertEqual(outbounds[0].content, "नमस्ते! मैं आपकी कैसे मदद कर सकता हूँ?")

    def test_profile_state_included_in_llm_context(self):
        apply_fields(self.customer, {"full_name": "Ravi", "nationality": "Indian"})
        captured = {}

        class FakeAdapter:
            def send(self, messages, config):
                captured["messages"] = messages
                return "ok"

        from llm.models import LLMConfig
        LLMConfig.objects.create(
            name="Test", provider="anthropic", api_key="k",
            model_name="claude-sonnet-4-6", system_prompt="Be helpful.", is_active=True)
        with mock.patch("bot.engine.get_adapter", return_value=FakeAdapter()):
            handle_inbound_message(self.conversation, "what trips do you offer?")
        system_text = " ".join(
            m["content"] for m in captured["messages"] if m["role"] == "system")
        self.assertIn("Travel profile", system_text)
        self.assertIn("captured: Name: Ravi", system_text)
        self.assertIn("still missing", system_text)

    def test_profile_context_names_exactly_one_next_detail(self):
        """The context must name ONE detail to ask next, in collection-
        priority order — the production bot stacked several questions or
        re-asked details the visitor had already given."""
        apply_fields(self.customer, {"full_name": "Ravi", "nationality": "Indian"})
        captured = {}

        class FakeAdapter:
            def send(self, messages, config):
                captured["messages"] = messages
                return "ok"

        LLMConfig.objects.create(
            name="Test", provider="anthropic", api_key="k",
            model_name="claude-sonnet-4-6", system_prompt="Be helpful.", is_active=True)
        with mock.patch("bot.engine.get_adapter", return_value=FakeAdapter()):
            handle_inbound_message(self.conversation, "what trips do you offer?")
        system_text = " ".join(
            m["content"] for m in captured["messages"] if m["role"] == "system")
        found = re.search(r"NEXT DETAIL TO ASK.*?: (.+?)\.", system_text)
        self.assertTrue(found, "profile context must name the next detail to ask")
        self.assertEqual(found.group(1), "WhatsApp number")

    def test_next_missing_detail_follows_collection_priority(self):
        """WhatsApp first (contactability), then the trip-defining facts, then
        identity extras — deliberately not the REQUIRED_FIELDS order."""
        profile = apply_fields(self.customer, {"full_name": "Ravi", "nationality": "Indian"})
        self.assertEqual(profile.next_missing_detail(), "WhatsApp number")
        apply_fields(self.customer, {"whatsapp_number": "+911234567890"})
        self.assertEqual(profile.next_missing_detail(), "Travel date")
        apply_fields(self.customer, {
            "travel_date": dt.date(2026, 12, 1), "trip_days": 7, "adults": 2,
            "residence_country": "India",
        })
        self.assertIsNone(profile.next_missing_detail())

    def test_final_check_is_the_last_system_message(self):
        """glm-style models attend to the END of the prompt, not the top: the
        FINAL CHECK checklist must ride immediately before the visitor's
        message, which itself stays last."""
        captured = []

        class FakeAdapter:
            def send(self, messages, config):
                captured.append(messages)
                return "ok"

        LLMConfig.objects.create(
            name="Test", provider="anthropic", api_key="k",
            model_name="claude-sonnet-4-6", system_prompt="Be helpful.", is_active=True)
        with mock.patch("bot.engine.get_adapter", return_value=FakeAdapter()):
            handle_inbound_message(self.conversation, "dubai 3 days please")

        messages = captured[0]
        self.assertEqual(messages[-1]["role"], "user")
        self.assertIn("dubai 3 days please", messages[-1]["content"])
        self.assertEqual(messages[-2]["role"], "system")
        self.assertIn("FINAL CHECK", messages[-2]["content"])

    def test_engine_enqueues_extraction_task(self):
        with mock.patch("profiles.tasks.extract_profile_task") as fake_task:
            handle_inbound_message(self.conversation, "hello there team")
        fake_task.delay.assert_called_once()
        args = fake_task.delay.call_args[0]
        self.assertEqual(args[0], str(self.customer.pk))
        self.assertEqual(args[1], "hello there team")


class WidgetSendTests(StubbedLLMMixin, TestCase):
    """The widget send flow with the async pipeline: /send/ answers with the
    typing-poller fragment and the bubbles arrive via /poll/ — these tests
    follow the poller exactly like the widget's JS does."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.channel = Channel.objects.create(
            name="Site", channel_type="wordpress", is_active=True,
            credentials={"site_key": "sk-prof"})

    def _send(self, message):
        return self.client.post(
            reverse("widget-send", args=["visitor1"]),
            {"site_key": "sk-prof", "message": message})

    def _poll(self, fragment: str) -> str:
        """Follow the poller's hx-get URL. The template writes `&amp;` inside
        the attribute; a real browser's HTML parser decodes it back to `&`
        before HTMX issues the GET — mirror that here."""
        match = re.search(r'hx-get="([^"]+)"', fragment)
        self.assertIsNotNone(match, "send response must contain the poller fragment")
        return self.client.get(html.unescape(match.group(1))).content.decode()

    def test_plain_greeting_renders_single_bot_bubble(self):
        resp = self._send("hi")
        body = resp.content.decode()
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn("Optimistic UI", body)
        self.assertNotIn("#}", body)
        poll_body = self._poll(body)
        # No travel intent → one reply bubble, no interrogation.
        self.assertEqual(poll_body.count('class="msg-row bot"'), 1)
        self.assertNotIn("WhatsApp number", poll_body)

    def test_travel_request_adds_scripted_questions_once(self):
        self._send("hi")
        poll_body = self._poll(self._send("I need a Dubai package").content.decode())
        self.assertEqual(poll_body.count('class="msg-row bot"'), 2)
        self.assertIn("travelling", poll_body)
        poll_body = self._poll(self._send("a 5 star hotel please").content.decode())
        self.assertEqual(poll_body.count('class="msg-row bot"'), 1)


class DashboardProfileTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.channel = Channel.objects.create(name="WP", channel_type="wordpress")
        cls.staff = User.objects.create_user("staff", password="x", is_staff=True)

    def setUp(self):
        self.client = Client(HTTP_HOST="localhost")
        self.client.force_login(self.staff)

    def _make_profile(self, external_id="lead-1", **overrides):
        customer = Customer.objects.create(channel=self.channel, external_id=external_id)
        conversation = Conversation.objects.create(
            customer=customer, last_message_at=timezone.now())
        Message.objects.create(conversation=conversation,
                               sender_type=Message.SenderType.CUSTOMER,
                               content="Hi, I want a 7 days Dubai package")
        Message.objects.create(conversation=conversation,
                               sender_type=Message.SenderType.BOT,
                               content="Sure! May I have your name and WhatsApp number?")
        fields = {
            "full_name": "Ravi Kumar", "whatsapp_number": "+911234567890",
            "nationality": "Indian", "residence_country": "India",
            "gcc_residence_card": True, "residence_card_expiry": dt.date(2027, 1, 31),
            "travel_date": dt.date(2026, 12, 1), "trip_days": 7, "adults": 2,
            "children_ages": "5, 8",
        }
        fields.update(overrides)
        return apply_fields(customer, fields)

    def test_list_shows_profiles(self):
        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profiles"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, profile.profile_number)
        self.assertContains(resp, "Ravi Kumar")

    def test_filters_by_completeness_and_query(self):
        complete = self._make_profile(external_id="lead-a")
        partial = self._make_profile(
            external_id="lead-b", full_name="Partial Pete", whatsapp_number="",
            nationality="", residence_country="", travel_date=None,
            trip_days=None, adults=None, children_ages="")
        resp = self.client.get(reverse("dashboard-profiles"), {"complete": "yes"})
        self.assertContains(resp, complete.profile_number)
        self.assertNotContains(resp, partial.profile_number)
        resp = self.client.get(reverse("dashboard-profiles"), {"complete": "no"})
        self.assertContains(resp, partial.profile_number)
        self.assertNotContains(resp, complete.profile_number)
        resp = self.client.get(reverse("dashboard-profiles"), {"q": "Ravi"})
        self.assertContains(resp, complete.profile_number)
        self.assertNotContains(resp, partial.profile_number)

    def test_csv_export(self):
        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profile-export-csv"))
        self.assertEqual(resp["Content-Type"], "text/csv")
        self.assertIn("attachment", resp["Content-Disposition"])
        body = resp.content.decode()
        self.assertIn("Profile Number", body)
        self.assertIn(profile.profile_number, body)
        self.assertIn("Ravi Kumar", body)

    def test_xlsx_export(self):
        from openpyxl import load_workbook

        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profile-export-xlsx"))
        self.assertIn("spreadsheetml", resp["Content-Type"])
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Profile Number")
        self.assertEqual(ws.cell(row=2, column=1).value, profile.profile_number)
        self.assertEqual(ws.cell(row=2, column=2).value, "Ravi Kumar")

    def test_card_and_transcript_png_downloads(self):
        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profile-card", args=[profile.pk]))
        self.assertEqual(resp["Content-Type"], "image/png")
        self.assertTrue(resp.content.startswith(b"\x89PNG"))
        self.assertIn("BP-000001_Ravi_Kumar.png", resp["Content-Disposition"])
        resp = self.client.get(reverse("dashboard-profile-transcript", args=[profile.pk]))
        self.assertEqual(resp["Content-Type"], "image/png")
        self.assertTrue(resp.content.startswith(b"\x89PNG"))
        self.assertIn("BP-000001_Ravi_Kumar_transcript.png", resp["Content-Disposition"])

    def test_full_report_png_download(self):
        """The combined report: profile card on top, full transcript below."""
        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profile-report", args=[profile.pk]))
        self.assertEqual(resp["Content-Type"], "image/png")
        self.assertTrue(resp.content.startswith(b"\x89PNG"))
        self.assertIn("BP-000001_Ravi_Kumar_report.png", resp["Content-Disposition"])

    def test_detail_page_is_read_only(self):
        """No manual entry: the page only displays what the AI extracted."""
        profile = self._make_profile()
        resp = self.client.get(reverse("dashboard-profile-detail", args=[profile.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ravi Kumar")
        self.assertContains(resp, "Profile card (PNG)")
        self.assertContains(resp, "Full report (PNG)")
        self.assertContains(resp, "Collected automatically by the AI")
        self.assertNotContains(resp, "Save changes")
        # The edit form rendered inputs for every profile field; none may exist
        # now. (base.html's logout <form> is fine — we target the field inputs.)
        self.assertNotContains(resp, 'name="full_name"')
        self.assertNotContains(resp, 'name="whatsapp_number"')
        # POSTing edits must be rejected and change nothing.
        resp = self.client.post(reverse("dashboard-profile-detail", args=[profile.pk]), {
            "full_name": "Manual Override",
        })
        self.assertEqual(resp.status_code, 405)
        profile.refresh_from_db()
        self.assertEqual(profile.full_name, "Ravi Kumar")

    def test_profile_pages_require_staff(self):
        profile = self._make_profile()
        anonymous = Client(HTTP_HOST="localhost")
        for url in (
            reverse("dashboard-profiles"),
            reverse("dashboard-profile-detail", args=[profile.pk]),
            reverse("dashboard-profile-card", args=[profile.pk]),
            reverse("dashboard-profile-transcript", args=[profile.pk]),
            reverse("dashboard-profile-report", args=[profile.pk]),
            reverse("dashboard-profile-export-csv"),
            reverse("dashboard-profile-export-xlsx"),
        ):
            self.assertEqual(anonymous.get(url).status_code, 302, url)